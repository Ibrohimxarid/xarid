"""Kontaktlar (kontaktlar.xlsx) va rezolyutsiyadagi ijrochilarni ularga moslash."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

HEADERS = [
    "Ism yoki tashkilot (rezolyutsiyadagi kabi)",
    "Email (bir nechta bo'lsa — vergul bilan)",
    "Boshqa yozilishlar (vergul bilan)",
    "Izoh",
]

_CYR_TO_LAT = {
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "yo", "ж": "j", "з": "z", "и": "i",
    "й": "y", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t",
    "у": "u", "ф": "f", "х": "x", "ц": "ts", "ч": "ch", "ш": "sh", "щ": "sh", "ъ": "", "ь": "", "ы": "i",
    "э": "e", "ю": "yu", "я": "ya", "ў": "o", "қ": "q", "ғ": "g", "ҳ": "h",
}
_EMAIL_RE = re.compile(r"^[^\s@,;]+@[^\s@,;]+\.[^\s@,;]+$")


def is_valid_email(value: str) -> bool:
    return bool(_EMAIL_RE.match(str(value or "").strip()))


def split_emails(value: str) -> list[str]:
    out: list[str] = []
    for part in re.split(r"[\s,;]+", str(value or "")):
        part = part.strip().lower()
        if is_valid_email(part) and part not in out:
            out.append(part)
    return out


def split_list(value: str) -> list[str]:
    return [x.strip() for x in re.split(r"[,;\n]+", str(value or "")) if x.strip()]


def name_tokens(value: str) -> list[str]:
    """"U.Xoshimova", "Xoshimova U." va "У.Хошимова" bir xil natija beradi."""
    latin = "".join(_CYR_TO_LAT.get(ch, ch) for ch in str(value or "").lower())
    latin = re.sub(r"[ʻʼ'‘’`´]", "", latin)
    return sorted(t for t in re.split(r"[^a-z0-9]+", latin) if t)


@dataclass
class Contact:
    name: str
    emails: list[str]
    aliases: list[str] = field(default_factory=list)


def match_executors(names: list[str], contacts: list[Contact]) -> tuple[list[str], list[str]]:
    """Natija: (emaillar, topilmagan ismlar). Faqat bitta kontaktga mos kelsa qabul qilinadi."""
    index = [(c, [k for k in (name_tokens(x) for x in [c.name, *c.aliases]) if k]) for c in contacts]
    emails: list[str] = []
    unmatched: list[str] = []
    for name in names:
        tokens = name_tokens(name)
        if not tokens:
            continue
        hits = [c for c, keys in index if tokens in keys]
        if not hits:
            # Qisman moslik: masalan, faqat familiya yoki qo'shimcha so'z ("TTPU rektori").
            def partial(keys: list[list[str]]) -> bool:
                return any(
                    (set(tokens) <= set(k) and max(map(len, tokens)) >= 4)
                    or (set(k) <= set(tokens) and max(map(len, k)) >= 4)
                    for k in keys)
            hits = [c for c, keys in index if partial(keys)]
        if len(hits) == 1:
            for e in hits[0].emails:
                if e not in emails:
                    emails.append(e)
        else:
            unmatched.append(name.strip() + (" (bir nechta kontaktga mos)" if hits else ""))
    return emails, unmatched


class ContactBook:
    """kontaktlar.xlsx — fayl o'zgarsa avtomatik qayta o'qiladi; bo'lmasa namuna yaratiladi."""

    def __init__(self, path: Path):
        self.path = Path(path)
        self._mtime: float | None = None
        self._contacts: list[Contact] = []
        if not self.path.exists():
            self._create_template()

    def _create_template(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Kontaktlar"
        ws.append(HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col, width in zip("ABCD", (40, 45, 45, 30)):
            ws.column_dimensions[col].width = width
        wb.save(self.path)

    def contacts(self) -> list[Contact]:
        try:
            mtime = self.path.stat().st_mtime
        except OSError:
            return []
        if mtime != self._mtime:
            self._contacts = self._read()
            self._mtime = mtime
        return self._contacts

    def _read(self) -> list[Contact]:
        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        try:
            out = []
            for row in wb.worksheets[0].iter_rows(min_row=2, max_col=3, values_only=True):
                name, emails, aliases = (str(v).strip() if v is not None else "" for v in (list(row) + [None] * 3)[:3])
                parsed = split_emails(emails)
                if name and parsed:
                    out.append(Contact(name, parsed, split_list(aliases)))
            return out
        finally:
            wb.close()

    def match(self, names: list[str]) -> tuple[list[str], list[str]]:
        return match_executors(names, self.contacts())
