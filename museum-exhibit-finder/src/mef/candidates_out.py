"""Russian-language outputs of the TPM candidate list: Markdown and Excel."""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

from .candidates import Candidate, Channel, build
from .config import load_yaml
from .evidence import parse_date
from .models import Evidence, MuseumFile
from .priority import _expired
from .relevance import CHECK_NAMES, _v, rules

YES_RU = {"yes": "да", "no": "нет", "unknown": "неизвестно"}
INTERNATIONAL_RU = {
    "yes": "да — ограничений для иностранных получателей нет",
    "domestic_first": "сначала организации своей страны, потом другие",
    "no": "нет — только организации своей страны",
    "unknown": "не указано",
}
OBJECT_TYPE_RU = {
    "interactive_station": "интерактивная станция", "demonstrator": "демонстрационная установка",
    "cutaway": "разрезная модель", "simulator": "симулятор", "robot": "робот",
    "digital_installation": "цифровая инсталляция", "engine": "двигатель", "vehicle": "транспортное средство",
    "aircraft": "самолёт / секция самолёта", "rolling_stock": "подвижной состав", "machine": "машина / станок",
    "instrument": "прибор", "computer": "компьютер", "model": "макет", "component": "деталь",
    "household": "бытовая техника", "display_furniture": "витрины / мебель", "book_document": "книги / документы",
    "decorative": "декор / сувениры", "mixed_collection": "партия без перечня", "unspecified": "не указан",
}
VERIFIED_VIA = {
    "web_search_snippet": "сниппет поисковой выдачи — открыть ссылку",
    "page_fetch": "страница загружена",
    "manual": "проверено вручную",
}


def _cell(v) -> str:
    return re.sub(r"\s+", " ", str(v or "")).replace("|", "\\|").strip()


def offer_status(evs: list[Evidence], today: dt.date, notes: list[str] | None = None) -> str:
    from .relevance import _notes_ru

    social = [n for n in (notes or []) if n.startswith("Availability reported only on social media")]
    if evs and social:
        return "сведения только из соцсетей — подтвердить на официальном сайте"
    if not evs:
        return "не предлагается — снят / на складе, судьба неизвестна"
    open_until = [e for e in evs if e.valid_until and parse_date(e.valid_until) >= today]
    if open_until:
        return f"открыто до {max(e.valid_until for e in open_until)}"
    closed = [e for e in evs if _expired(e, today)]
    if closed and len(closed) == len(evs):
        return f"срок истёк {max(e.valid_until for e in closed)} — спросить, что осталось"
    dated = [e for e in evs if e.source_date]
    if dated:
        return f"предложение от {max(e.source_date for e in dated)} — уточнить, что ещё доступно"
    return "дата не указана — проверить, что объявление активно"


def _ev_lines(evs: list[Evidence], today: dt.date) -> str:
    out = []
    for e in evs:
        closed = f" [срок истёк {e.valid_until}]" if _expired(e, today) else ""
        out.append(f"[{e.id}] {e.claim.rstrip('.')}.{closed} ({e.source_date or 'без даты'}) {e.url}")
    return "\n".join(out)


def _contacts(c: Candidate | Channel) -> dict[str, str]:
    from .rows import ranked_contacts

    cs = ranked_contacts(c.museum)
    join = lambda vals: "\n".join(dict.fromkeys(v for v in vals if v))  # noqa: E731
    return {
        "Контактное лицо / отдел": join(", ".join(x for x in [k.name, k.position] if x) for k in cs),
        "Email": join(k.email for k in cs),
        "Телефон": join(k.phone for k in cs),
        "Контактная страница": join(k.contact_page for k in cs),
    }


CANDIDATE_FIELDS = [
    "№", "ACQUISITION_PRIORITY", "MUSEUM_RELEVANCE", "Экспонат", "Тип объекта", "Направление профиля TPM",
    "WHY_IT_FITS_TASHKENT_POLYTECHNIC_MUSEUM", "Музей", "Страна", "История", "Статус предложения",
    "Почему не HIGH (открытые проверки)", *CHECK_NAMES[1:], "Фото / документация", "Иностранный получатель",
    "Кто может получить", "Доказательство: предлагают", "Доказательство: снят / заменён / закрыт",
    "Ссылка (главный источник)", "Дата источника", "Контактное лицо / отдел", "Email", "Телефон",
    "Контактная страница", "Следующий шаг", "Кратко (рус.)", "Проверено через",
]
CHANNEL_FIELDS = [
    "Музей / программа", "Страна", "Что известно о составе", "Доказательство", "Ссылка", "Дата",
    "Иностранный получатель", "Кто может получить", "Email", "Телефон", "Контактная страница",
    "Следующий шаг", "Кратко (рус.)",
]
REJECTED_FIELDS = ["Музей", "Страна", "Объект", "Тип объекта", "Решение", "Причина", "Ссылка"]


def candidate_record(n: int, c: Candidate, today: dt.date) -> dict:
    a, acq, mf = c.assessment, c.acquisition, c.museum
    offer = c.offer_evidence()
    pool = offer or c.removal_evidence() or c.evidence
    primary = max(pool, key=lambda e: parse_date(e.source_date) or dt.date.min) if pool else None
    checks = {ch.name: f"{YES_RU[ch.status]} — {ch.detail}" for ch in acq.checks}
    why = c.why_it_fits
    if c.why_problems:
        why = (why + "\n" if why else "") + "⚠ " + "; ".join(c.why_problems)
    return {
        "№": n,
        "ACQUISITION_PRIORITY": acq.priority,
        "MUSEUM_RELEVANCE": acq.relevance.value,
        "Экспонат": c.label + ("" if c.exhibit else " [галерея, предметы не перечислены]"),
        "Тип объекта": OBJECT_TYPE_RU.get(_v(c.obj.object_type), _v(c.obj.object_type)),
        "Направление профиля TPM": c.areas,
        "WHY_IT_FITS_TASHKENT_POLYTECHNIC_MUSEUM": why,
        "Музей": mf.museum.name,
        "Страна": mf.museum.country,
        "История": c.history(),
        "Статус предложения": offer_status(offer, today, a.notes),
        "Почему не HIGH (открытые проверки)": acq.why_not_high() if acq.priority != "HIGH" else "—",
        **{name: checks.get(name, "") for name in CHECK_NAMES[1:]},
        "Фото / документация": c.ideal()["photos_docs"],
        "Иностранный получатель": INTERNATIONAL_RU[_v(mf.research.international_transfer)],
        "Кто может получить": mf.research.eligibility or "не указано — уточнить",
        "Доказательство: предлагают": _ev_lines(offer, today) or "—",
        "Доказательство: снят / заменён / закрыт": _ev_lines(c.removal_evidence(), today) or "—",
        "Ссылка (главный источник)": primary.url if primary else "",
        "Дата источника": (primary.source_date or "без даты") if primary else "",
        **_contacts(c),
        "Следующий шаг": mf.research.next_step or "",
        "Кратко (рус.)": mf.research.summary_ru or "",
        "Проверено через": ", ".join(sorted({VERIFIED_VIA.get(e.verified_via, e.verified_via) for e in c.evidence})),
    }


def channel_record(ch: Channel, today: dt.date, has_candidates: bool = False) -> dict:
    mf = ch.museum
    evs = ch.offer_evidence or [e for e in mf.evidence if set(e.signals) - {"contact"}]
    primary = evs[0] if evs else None
    items = "; ".join(ch.items) if ch.items else (
        "отдельные объекты — в листе «Кандидаты»; полный список запросить" if has_candidates
        else "список не опубликован / не получен")
    name = mf.museum.name + (f" — {ch.programme}" if ch.programme else "")
    contacts = _contacts(ch)
    return {
        "Музей / программа": name,
        "Страна": mf.museum.country,
        "Что известно о составе": items,
        "Доказательство": _ev_lines(evs, today) or "—",
        "Ссылка": primary.url if primary else "",
        "Дата": (primary.source_date or "без даты") if primary else "",
        "Иностранный получатель": INTERNATIONAL_RU[_v(mf.research.international_transfer)],
        "Кто может получить": mf.research.eligibility or "не указано — уточнить",
        "Email": contacts["Email"],
        "Телефон": contacts["Телефон"],
        "Контактная страница": contacts["Контактная страница"],
        "Следующий шаг": mf.research.next_step or "",
        "Кратко (рус.)": mf.research.summary_ru or "",
    }


def rejected_record(c: Candidate) -> dict:
    primary = (c.offer_evidence() or c.evidence or [None])[0]
    return {
        "Музей": c.museum.museum.name,
        "Страна": c.museum.museum.country,
        "Объект": c.label,
        "Тип объекта": OBJECT_TYPE_RU.get(_v(c.obj.object_type), _v(c.obj.object_type)),
        "Решение": "НЕ РЕКОМЕНДОВАТЬ (NOT_RELEVANT)",
        "Причина": c.acquisition.relevance_reason + " — не соответствует профилю TPM "
                   "(наука, техника, инженерия, транспорт).",
        "Ссылка": primary.url if primary else "",
    }


# ---------------------------------------------------------------- Markdown

def candidates_markdown(files: list[MuseumFile], today: dt.date | None = None) -> str:
    today = today or dt.date.today()
    cands, chans, rejected = build(files, today)
    by_p = {p: [c for c in cands if c.acquisition.priority == p] for p in ("HIGH", "MEDIUM", "LOW")}
    offered_low = [c for c in by_p["LOW"] if c.offered]
    ask_low = [c for c in by_p["LOW"] if not c.offered]
    out = [
        "# Кандидаты для Ташкентского политехнического музея — проверка соответствия профилю",
        "",
        f"Дата: {today.isoformat()}. Каждый объект прошёл проверку соответствия профилю TPM "
        "(TASHKENT POLYTECHNIC MUSEUM RELEVANCE CHECK): автомобильная техника, транспорт, инженерия, "
        "механика, физика, электричество, оптика, робототехника, промышленные технологии, STEM, "
        "интерактивная наука, история техники.",
        "",
        "- **MUSEUM_RELEVANCE** определяется правилами по двум фактам: категория объекта и что он "
        "физически собой представляет (интерактивная станция, разрезная модель, двигатель, автомобиль, "
        "книга, витрина…). Без субъективных баллов. Правила: `config/relevance.yaml`.",
        "- **ACQUISITION_PRIORITY** — это не только доступность. Проверяются шесть фактов: сильное "
        "соответствие, подлинный экспонат, физически пригоден, образовательная ценность, разумная "
        "логистика (включая право иностранного музея получить объект) и подтверждённая доступность "
        "сейчас. **HIGH** — всё «да»; **MEDIUM** — одна открытая проверка; **LOW** — две и больше или "
        "слабое соответствие.",
        "- Объекты NOT_RELEVANT (книги, мебель, витрины, декор) в список не попадают — они перечислены "
        "в конце как «не рекомендовать».",
        "",
        "> ⚠️ Доказательства взяты из сниппетов поисковой выдачи (страницы не открывались из среды "
        "исследования). Перед письмом откройте ссылку.",
        "",
        "## Итог",
        "",
        f"| Приоритет | Объектов |\n|---|---|\n| HIGH | {len(by_p['HIGH'])} |\n| MEDIUM | {len(by_p['MEDIUM'])} |\n"
        f"| LOW — предлагали (срок истёк / есть препятствия) | {len(offered_low)} |\n"
        f"| LOW — сняты, судьба неизвестна (спросить) | {len(ask_low)} |\n"
        f"| Каналы передачи (список предметов запросить) | {len(chans)} |\n"
        f"| Не рекомендовать (NOT_RELEVANT) | {len(rejected)} |",
        "",
    ]
    if not by_p["HIGH"]:
        out += ["**Ни один объект пока не проходит все проверки HIGH.** Причины по каждому объекту — в "
                "колонке «Почему не HIGH».", ""]

    def table(rows: list[Candidate], start: int = 1) -> list[str]:
        t = ["| № | Приоритет | Соответствие | Экспонат | Музей | Статус предложения | Почему не HIGH |",
             "|---|---|---|---|---|---|---|"]
        for i, c in enumerate(rows, start):
            acq = c.acquisition
            t.append(f"| {i} | {acq.priority} | {acq.relevance.value} | {_cell(c.label)} | "
                     f"{_cell(c.museum.museum.name)} ({_cell(c.museum.museum.country)}) | "
                     f"{_cell(offer_status(c.offer_evidence(), today, c.assessment.notes))} | "
                     f"{_cell(acq.why_not_high() if acq.priority != 'HIGH' else '—')} |")
        return t + [""]

    def cards(rows: list[Candidate], start: int = 1) -> list[str]:
        t = []
        for i, c in enumerate(rows, start):
            r = candidate_record(i, c, today)
            t += [f"### {i}. {c.label} — {c.museum.museum.name} ({c.museum.museum.country})", ""]
            t += [f"- **ACQUISITION_PRIORITY:** {r['ACQUISITION_PRIORITY']} · **MUSEUM_RELEVANCE:** "
                  f"{r['MUSEUM_RELEVANCE']} ({_cell(c.acquisition.relevance_reason)})",
                  f"- **Тип объекта:** {r['Тип объекта']} · **Направление:** {r['Направление профиля TPM']}",
                  f"- **WHY_IT_FITS_TASHKENT_POLYTECHNIC_MUSEUM:** {_cell(r['WHY_IT_FITS_TASHKENT_POLYTECHNIC_MUSEUM'])}",
                  f"- **История:** {_cell(r['История'])}",
                  f"- **Статус предложения:** {r['Статус предложения']}",
                  f"- **Проверки:** " + "; ".join(f"{ch.name} — {YES_RU[ch.status]}" for ch in c.acquisition.checks[1:])]
            if c.acquisition.priority != "HIGH":
                t.append(f"- **Почему не HIGH:** {_cell(r['Почему не HIGH (открытые проверки)'])}")
            t += [f"- **Иностранный получатель:** {r['Иностранный получатель']}",
                  f"- **Кто может получить:** {_cell(r['Кто может получить'])}"]
            if c.offer_evidence():
                t.append("- **Доказательство предложения:**")
                t += [f"  - {_cell(x)}" for x in r["Доказательство: предлагают"].split("\n")]
            if c.removal_evidence():
                t.append("- **Доказательство снятия / закрытия:**")
                t += [f"  - {_cell(x)}" for x in r["Доказательство: снят / заменён / закрыт"].split("\n")]
            contact = "; ".join(x for x in [r["Контактное лицо / отдел"], r["Email"], r["Телефон"],
                                            r["Контактная страница"]] if x).replace("\n", ", ")
            t += [f"- **Контакт:** {contact or 'не найден'}",
                  f"- **Следующий шаг:** {_cell(r['Следующий шаг']) or '—'}", ""]
        return t

    n = 1
    for title, rows in (
        ("## HIGH — сильное соответствие, всё подтверждено", by_p["HIGH"]),
        ("## MEDIUM — сильное соответствие, одна открытая проверка", by_p["MEDIUM"]),
        ("## LOW — предлагали, но есть два и более препятствия", offered_low),
    ):
        out += [title, ""]
        if rows:
            out += table(rows, n) + cards(rows, n)
            n += len(rows)
        else:
            out += ["_Нет._", ""]

    out += ["## LOW — подходящие экспонаты сняты, судьба неизвестна (спросить у музея)", "",
            "Здесь только объекты и галереи с сильным или прямым соответствием профилю, которые по источнику "
            "сняты или заменены, но предложения о передаче нет.", ""]
    strong_ask = [c for c in ask_low if c.acquisition.relevance.value in ("DIRECT_MATCH", "STRONG_MATCH")]
    out += table(strong_ask, n) if strong_ask else ["_Нет._", ""]
    n += len(strong_ask)

    cand_museums = {c.museum.museum.id for c in cands}
    out += ["## Каналы: музей реально передаёт предметы, но список не получен", "",
            "| Музей / программа | Что известно | Иностранный получатель | Источник |", "|---|---|---|---|"]
    for ch in chans:
        r = channel_record(ch, today, ch.museum.museum.id in cand_museums)
        out.append(f"| {_cell(r['Музей / программа'])} ({_cell(r['Страна'])}) | {_cell(r['Что известно о составе'])} | "
                   f"{_cell(r['Иностранный получатель'])} | {r['Ссылка']} ({r['Дата']}) |")
    out.append("")

    out += ["## Не рекомендовать (NOT_RELEVANT)", "",
            "Эти предметы реально предлагаются, но не подходят профилю TPM — в работу не брать.", "",
            "| Музей | Объект | Причина |", "|---|---|---|"]
    out += [f"| {_cell(c.museum.museum.name)} | {_cell(c.label)} | {_cell(c.acquisition.relevance_reason)} |"
            for c in rejected] or ["| — | — | — |"]
    out.append("")
    excluded = load_yaml("excluded_channels.yaml").get("excluded", [])
    if excluded:
        out += ["## Исключённые каналы (продажа новых выставок / только для своей страны)", "",
                "| Канал | Почему | Источник |", "|---|---|---|"]
        out += [f"| {_cell(x['name'])} | {_cell(x['reason'])} | {x['url']} |" for x in excluded]
        out.append("")
    return "\n".join(out)


# ---------------------------------------------------------------- Excel

WIDTHS = {
    "№": 5, "ACQUISITION_PRIORITY": 14, "MUSEUM_RELEVANCE": 16, "Экспонат": 40, "Тип объекта": 18,
    "Направление профиля TPM": 28, "WHY_IT_FITS_TASHKENT_POLYTECHNIC_MUSEUM": 60, "Музей": 28, "Страна": 14,
    "История": 36, "Статус предложения": 30, "Почему не HIGH (открытые проверки)": 60,
    "Фото / документация": 14, "Иностранный получатель": 30, "Кто может получить": 40,
    "Доказательство: предлагают": 70, "Доказательство: снят / заменён / закрыт": 60,
    "Ссылка (главный источник)": 40, "Дата источника": 12, "Контактное лицо / отдел": 28, "Email": 30,
    "Телефон": 18, "Контактная страница": 40, "Следующий шаг": 50, "Кратко (рус.)": 60, "Проверено через": 26,
    "Музей / программа": 45, "Что известно о составе": 45, "Доказательство": 70, "Ссылка": 45, "Дата": 12,
    "Объект": 45, "Решение": 26, "Причина": 70,
}
PRIORITY_FILL = {"HIGH": "C6EFCE", "MEDIUM": "FFEB9C", "LOW": "F2F2F2"}
RELEVANCE_FILL = {"DIRECT_MATCH": "A9D08E", "STRONG_MATCH": "C6E0B4", "RELATED": "FFF2CC", "WEAK_MATCH": "FCE4D6"}


def _sheet(wb, title: str, fields: list[str], records: list[dict], fill: str):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(title)
    ws.append(fields)
    for r in records:
        ws.append([r.get(f, "") for f in fields])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 42
    links = {i for i, f in enumerate(fields) if f.startswith("Ссылка") or f == "Контактная страница"}
    for row in ws.iter_rows(min_row=2):
        for i, cell in enumerate(row):
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            v = str(cell.value or "")
            if i in links and v.startswith("http") and "\n" not in v:
                cell.hyperlink = v
                cell.font = Font(color="0563C1", underline="single")
            if fields[i] == "ACQUISITION_PRIORITY" and v in PRIORITY_FILL:
                cell.fill = PatternFill("solid", fgColor=PRIORITY_FILL[v])
                cell.font = Font(bold=True)
            if fields[i] == "MUSEUM_RELEVANCE" and v in RELEVANCE_FILL:
                cell.fill = PatternFill("solid", fgColor=RELEVANCE_FILL[v])
    for i, f in enumerate(fields, start=1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(f, 22)
    ws.freeze_panes = "E2" if "Экспонат" in fields else "B2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def write_candidates_xlsx(files: list[MuseumFile], path: Path, today: dt.date | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    today = today or dt.date.today()
    cands, chans, rejected = build(files, today)
    offered = [c for c in cands if c.offered or c.acquisition.priority in ("HIGH", "MEDIUM")]
    ask = [c for c in cands if c not in offered]
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Кандидаты", CANDIDATE_FIELDS, [candidate_record(i, c, today) for i, c in enumerate(offered, 1)],
           "1E6B3A")
    _sheet(wb, "Сняты — спросить", CANDIDATE_FIELDS,
           [candidate_record(i, c, today) for i, c in enumerate(ask, 1)], "9C6500")
    cand_museums = {c.museum.museum.id for c in cands}
    _sheet(wb, "Каналы — запросить список", CHANNEL_FIELDS,
           [channel_record(ch, today, ch.museum.museum.id in cand_museums) for ch in chans], "1F3A5F")
    _sheet(wb, "Не рекомендовать", REJECTED_FIELDS, [rejected_record(c) for c in rejected], "7F7F7F")
    excluded = [{"Канал": x["name"], "Почему исключён": x["reason"], "Ссылка": x["url"]}
                for x in load_yaml("excluded_channels.yaml").get("excluded", [])]
    _sheet(wb, "Исключённые каналы", ["Канал", "Почему исключён", "Ссылка"], excluded, "7F7F7F")

    ws = wb.create_sheet("Правила")
    r = rules()
    lines = [
        f"TASHKENT POLYTECHNIC MUSEUM RELEVANCE CHECK — {today.isoformat()}",
        "",
        "Профиль музея: " + ", ".join(r["museum_profile"]),
        "",
        "MUSEUM_RELEVANCE (категории, не баллы):",
        "DIRECT_MATCH — интерактивная станция, демонстрационная установка, разрезная модель, симулятор, робот, "
        "цифровая инсталляция или двигатель в основном направлении (A1–A9).",
        "STRONG_MATCH — реальный технический объект (автомобиль, самолёт, подвижной состав, машина, прибор, "
        "компьютер, макет) в основном направлении или в группе B.",
        "RELATED — деталь технического объекта.",
        "WEAK_MATCH — бытовая техника или слабая связь с профилем.",
        "NOT_RELEVANT — книги, документы, мебель, витрины, декор или нет связи с техникой. Не показывается "
        "в кандидатах.",
        "",
        "ACQUISITION_PRIORITY — шесть проверок: " + "; ".join(CHECK_NAMES) + ".",
        "HIGH — сильное соответствие (DIRECT/STRONG) и все остальные проверки «да».",
        "MEDIUM — сильное соответствие и ровно одна проверка «нет» или «неизвестно».",
        "LOW — слабое соответствие (RELATED/WEAK) или две и больше открытых проверок.",
        "«Разумная логистика» включает право иностранного музея получить объект: если источник ставит "
        "в приоритет организации своей страны или не говорит об иностранных получателях — «неизвестно».",
        "",
        "Направления профиля:",
        *[f"{k} — {v}" for k, v in r["areas"].items()],
        "",
        "Листы: «Кандидаты» — предлагаются или предлагались (и все HIGH/MEDIUM); «Сняты — спросить» — "
        "подходящие объекты сняты, но предложения нет; «Каналы» — музей реально отдаёт предметы, список "
        "не получен; «Не рекомендовать» — предлагаются, но NOT_RELEVANT.",
        "",
        "ВАЖНО: доказательства получены из сниппетов поисковой выдачи. Перед письмом откройте ссылку.",
    ]
    for line in lines:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 130
    for row in ws.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
