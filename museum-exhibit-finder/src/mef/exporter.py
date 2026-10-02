"""[9] CSV / Excel exporter."""

from __future__ import annotations

import csv
import datetime as dt
from pathlib import Path

from .config import PATHS
from .models import MuseumFile
from .priority import assess_museum
from .rows import DATABASE_FIELDS, database_record, iter_rows
from .stats import compute_stats
from .candidates_out import write_candidates_xlsx

MUSEUM_FIELDS = [
    "Museum ID", "Museum", "Local Name", "Country", "City", "Website", "Museum Type",
    "Parent Org", "General Email", "Phone", "LinkedIn", "Priority", "Chain Stage",
    "Availability", "Exhibitions Tracked", "Exhibits Tracked", "Evidence Count",
    "Outreach Status", "Next Step", "Last Updated",
]
EVIDENCE_FIELDS = [
    "Museum ID", "Museum", "Evidence ID", "URL", "Title", "Publisher", "Source Type",
    "Source Date", "Accessed", "Language", "Claim", "Excerpt", "Verbatim", "Signals", "Verified Via",
]
CONTACT_FIELDS = [
    "Museum ID", "Museum", "Country", "Name", "Position", "Kind", "Email", "Phone",
    "LinkedIn", "Contact Page", "Source",
]


def museum_records(files: list[MuseumFile]) -> list[dict]:
    out = []
    for mf in files:
        m, a = mf.museum, assess_museum(mf)
        out.append({
            "Museum ID": m.id, "Museum": m.name, "Local Name": m.name_local or "",
            "Country": m.country, "City": m.city or "", "Website": m.website or "",
            "Museum Type": m.type, "Parent Org": m.parent_org or "",
            "General Email": m.general_email or "", "Phone": m.phone or "",
            "LinkedIn": m.linkedin or "", "Priority": a.priority or "—", "Chain Stage": a.stage,
            "Availability": a.availability, "Exhibitions Tracked": len(mf.exhibitions),
            "Exhibits Tracked": sum(len(e.exhibits) for e in mf.exhibitions),
            "Evidence Count": len(mf.evidence), "Outreach Status": mf.research.outreach_status,
            "Next Step": mf.research.next_step or "", "Last Updated": mf.research.last_updated,
        })
    return out


def evidence_records(files: list[MuseumFile]) -> list[dict]:
    return [
        {
            "Museum ID": mf.museum.id, "Museum": mf.museum.name, "Evidence ID": e.id, "URL": e.url,
            "Title": e.title or "", "Publisher": e.publisher or "", "Source Type": e.source_type,
            "Source Date": e.source_date or "", "Accessed": e.accessed, "Language": e.language or "",
            "Claim": e.claim, "Excerpt": e.excerpt or "", "Verbatim": "yes" if e.excerpt_is_verbatim else "no",
            "Signals": ", ".join(e.signals), "Verified Via": e.verified_via,
        }
        for mf in files
        for e in mf.evidence
    ]


def contact_records(files: list[MuseumFile]) -> list[dict]:
    out = []
    for mf in files:
        by_id = mf.evidence_by_id()
        for c in mf.contacts:
            src = c.source if c.source.startswith("http") else (by_id[c.source].url if c.source in by_id else c.source)
            out.append({
                "Museum ID": mf.museum.id, "Museum": mf.museum.name, "Country": mf.museum.country,
                "Name": c.name or "", "Position": c.position or "", "Kind": c.kind,
                "Email": c.email or "", "Phone": c.phone or "", "LinkedIn": c.linkedin or "",
                "Contact Page": c.contact_page or "", "Source": src,
            })
    return out


def _write_csv(path: Path, fields: list[str], records: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(records)


def _sheet(wb, title: str, fields: list[str], records: list[dict]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(title)
    ws.append(fields)
    for r in records:
        ws.append([r.get(f, "") for f in fields])
    head = PatternFill("solid", fgColor="1F3A5F")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    fills = {"A": "C6EFCE", "B": "FFEB9C", "C": "DDEBF7"}
    if "Priority" in fields:
        col = fields.index("Priority") + 1
        for row in ws.iter_rows(min_row=2):
            p = str(row[col - 1].value or "")
            if p in fills:
                row[col - 1].fill = PatternFill("solid", fgColor=fills[p])
    for i, f in enumerate(fields, start=1):
        width = max([len(f)] + [min(60, len(str(r.get(f, "")))) for r in records[:200]])
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(60, width + 2))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def export_all(
    files: list[MuseumFile], out_dir: Path | None = None, contacts_dir: Path | None = None
) -> dict[str, Path]:
    out_dir = out_dir or PATHS.exports
    contacts_dir = contacts_dir or PATHS.contacts
    rows = iter_rows(files)
    db_records = [database_record(r) for r in rows]
    order = {"A": 0, "B": 1, "C": 2, "—": 3}
    db_records.sort(key=lambda r: (order.get(r["Priority"], 4), r["Country"], r["Museum"]))
    # Exhibit opportunities: removed/offered exhibits that passed the TPM relevance check
    # (NOT_RELEVANT objects — books, furniture, decorative — are never shown here).
    opportunities = [r for r in db_records if r["Priority"] in ("A", "B") and r["Exhibit ID"]
                     and r["Museum Relevance"] in ("DIRECT_MATCH", "STRONG_MATCH", "RELATED", "WEAK_MATCH")]
    opportunities.sort(key=lambda r: ({"HIGH": 0, "MEDIUM": 1, "LOW": 2}.get(r["Acquisition Priority"], 3),
                                      r["Museum Relevance"], r["Museum"]))
    museums = sorted(museum_records(files), key=lambda r: (order.get(r["Priority"], 4), r["Country"], r["Museum"]))
    evidence = evidence_records(files)
    contacts = contact_records(files)
    stats = compute_stats(files)

    paths = {
        "database": out_dir / "database.csv",
        "opportunities": out_dir / "exhibit_opportunities.csv",
        "museums": out_dir / "museums.csv",
        "evidence": out_dir / "sources.csv",
        "contacts": out_dir / "contacts.csv",
        "contacts_dir": contacts_dir / "contacts.csv",
        "xlsx": out_dir / "museum_finder.xlsx",
        "candidates_xlsx": out_dir / "tpm_candidates.xlsx",
    }
    _write_csv(paths["database"], DATABASE_FIELDS, db_records)
    _write_csv(paths["opportunities"], DATABASE_FIELDS, opportunities)
    _write_csv(paths["museums"], MUSEUM_FIELDS, museums)
    _write_csv(paths["evidence"], EVIDENCE_FIELDS, evidence)
    _write_csv(paths["contacts"], CONTACT_FIELDS, contacts)
    _write_csv(paths["contacts_dir"], CONTACT_FIELDS, contacts)

    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Database", DATABASE_FIELDS, db_records)
    _sheet(wb, "Exhibit Opportunities", DATABASE_FIELDS, opportunities)
    _sheet(wb, "Museum Leads", MUSEUM_FIELDS, museums)
    _sheet(wb, "Contacts", CONTACT_FIELDS, contacts)
    _sheet(wb, "Sources", EVIDENCE_FIELDS, evidence)
    _sheet(wb, "Statistics", ["Metric", "Value"],
           [{"Metric": k, "Value": v} for k, v in stats.items()] +
           [{"Metric": "Generated", "Value": dt.date.today().isoformat()}])
    paths["xlsx"].parent.mkdir(parents=True, exist_ok=True)
    wb.save(paths["xlsx"])
    write_candidates_xlsx(files, paths["candidates_xlsx"])
    return paths
