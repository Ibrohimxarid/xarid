"""Excel version of the verified-offers list (``exports/verified_offers.xlsx``)."""

from __future__ import annotations

import datetime as dt
from pathlib import Path

from .config import load_yaml
from .models import Evidence, MuseumFile
from .config import settings
from .evidence import age_years
from .priority import UNCONFIRMED_SOURCES, _expired
from .rows import ranked_contacts
from .verified import _museum_evidence, _offer_status, collect


def _ev_text(evs: list[Evidence]) -> str:
    return "\n".join(f"[{e.id}] {e.claim.rstrip('.')}. ({e.source_date or 'без даты'}) {e.url}" for e in evs)


def _contacts(mf: MuseumFile) -> dict[str, str]:
    cs = ranked_contacts(mf)
    join = lambda vals: "\n".join(dict.fromkeys(v for v in vals if v))  # noqa: E731
    return {
        "Контактное лицо / отдел": join(", ".join(x for x in [c.name, c.position] if x) for c in cs),
        "Email": join(c.email for c in cs),
        "Телефон": join(c.phone for c in cs),
        "Контактная страница": join(c.contact_page for c in cs),
    }


VERIFIED_FIELDS = [
    "№", "Музей", "Страна", "Город", "Что предлагают", "Категория", "Размеры", "Кратко (рус.)",
    "Статус предложения", "Кто может получить", "Доказательство: предлагают / отдают / продают",
    "Ссылка на предложение", "Дата источника", "Доказательство: сняты / списаны / музей закрыт",
    "Более ранние раунды (срок истёк)", "Контактное лицо / отдел", "Email", "Телефон",
    "Контактная страница", "Следующий шаг", "Проверено через",
]
CLOSED_FIELDS = [
    "Музей", "Страна", "Что предлагали", "Кратко (рус.)", "Почему не в основном списке",
    "Ссылка", "Email", "Телефон", "Контактная страница", "Следующий шаг",
]
EXCLUDED_FIELDS = ["Канал", "Почему исключён", "Ссылка"]
EVIDENCE_FIELDS = [
    "Музей", "ID", "Роль", "Утверждение (из источника)", "Текст источника", "Дата источника",
    "Срок подачи до", "Ссылка", "Тип источника", "Проверено через",
]
CATEGORY_RU = {
    "acoustics": "акустика", "automotive": "автомобили", "aviation": "авиация", "computing": "вычислительная техника",
    "display_cases": "витрины", "electricity": "электричество", "engineering": "инженерия",
    "industrial": "промышленность", "mechanics": "механика", "optics": "оптика", "physics": "физика",
    "railway": "железная дорога", "simulator": "симуляторы", "space": "космос", "stem": "STEM",
    "telecommunications": "связь", "thermodynamics": "термодинамика", "transport": "транспорт",
}


def _reason_ru(evs: list[Evidence], today: dt.date) -> str:
    """Why an earlier offer is not in the main list, in Russian, derived from the evidence."""
    reasons = []
    closed = [e.valid_until for e in evs if _expired(e, today)]
    if closed:
        reasons.append(f"Срок подачи истёк {max(closed)}.")
    max_age = settings().availability_max_age_years
    old = [e.source_date for e in evs if not _expired(e, today) and (age_years(e, today) or 0) > max_age]
    if old:
        reasons.append(f"Данные от {max(old)} — старше {max_age} лет.")
    if evs and all(e.source_type in UNCONFIRMED_SOURCES for e in evs):
        reasons.append("Сведения только из соцсетей — нет официального подтверждения.")
    action = "Спросить, что осталось." if closed or old else "Подтвердить на официальном сайте музея."
    return " ".join(reasons + [action])


VERIFIED_VIA = {
    "web_search_snippet": "сниппет поисковой выдачи — открыть ссылку",
    "page_fetch": "страница загружена",
    "manual": "проверено вручную",
}


def _items_by_name(mf: MuseumFile) -> dict:
    by = {}
    for ex in mf.exhibitions:
        by[ex.name] = (ex, None)
        for it in ex.exhibits:
            by[it.name] = (ex, it)
    return by


def verified_records(files: list[MuseumFile], today: dt.date):
    verified, closed = collect(files, today)
    rows, evidence_rows = [], []
    for n, (mf, items) in enumerate(verified, 1):
        offer, earlier, removal = _museum_evidence(mf, items, today)
        by_name = _items_by_name(mf)
        for label, evs in items:
            _, it = next(((ex, it) for name, (ex, it) in by_name.items() if label.startswith(name)), (None, None))
            primary = evs[0] if evs else offer[0]
            rows.append({
                "№": n, "Музей": mf.museum.name, "Страна": mf.museum.country, "Город": mf.museum.city or "",
                "Что предлагают": label,
                "Категория": ", ".join(CATEGORY_RU.get(c, c) for c in it.category) if it else "",
                "Размеры": (it.dimensions or "") if it else "",
                "Кратко (рус.)": mf.research.summary_ru or "",
                "Статус предложения": _offer_status(offer, today),
                "Кто может получить": mf.research.eligibility or "не указано — уточнить",
                "Доказательство: предлагают / отдают / продают": _ev_text(offer),
                "Ссылка на предложение": primary.url,
                "Дата источника": primary.source_date or "без даты",
                "Доказательство: сняты / списаны / музей закрыт": _ev_text(removal),
                "Более ранние раунды (срок истёк)": _ev_text(earlier),
                **_contacts(mf),
                "Следующий шаг": mf.research.next_step or "",
                "Проверено через": ", ".join(sorted({VERIFIED_VIA.get(e.verified_via, e.verified_via) for e in offer})),
            })
        for role, evs in (("предложение", offer), ("ранний раунд (срок истёк)", earlier), ("снятие / закрытие", removal)):
            evidence_rows += [_evidence_row(mf, e, role, today) for e in evs]

    closed_rows = []
    for mf, earlier in closed:
        evs = list({e.id: e for _, a in earlier for e in a.availability_evidence}.values())
        closed_rows.append({
            "Музей": mf.museum.name, "Страна": mf.museum.country,
            "Что предлагали": "\n".join(label for label, _ in earlier),
            "Кратко (рус.)": mf.research.summary_ru or "",
            "Почему не в основном списке": _reason_ru(evs, today),
            "Ссылка": evs[0].url if evs else "",
            **{k: v for k, v in _contacts(mf).items() if k != "Контактное лицо / отдел"},
            "Следующий шаг": mf.research.next_step or "",
        })
        evidence_rows += [_evidence_row(mf, e, "ранее предлагали", today) for e in evs]

    excluded = [
        {"Канал": x["name"], "Почему исключён": x["reason"], "Ссылка": x["url"]}
        for x in load_yaml("excluded_channels.yaml").get("excluded", [])
    ]
    return rows, closed_rows, excluded, evidence_rows


def _evidence_row(mf: MuseumFile, e: Evidence, role: str, today: dt.date) -> dict:
    until = e.valid_until or ""
    if until and _expired(e, today):
        until += " (истёк)"
    return {
        "Музей": mf.museum.name, "ID": e.id, "Роль": role, "Утверждение (из источника)": e.claim,
        "Текст источника": (e.excerpt or "") + ("" if e.excerpt_is_verbatim or not e.excerpt else " [не дословно]"),
        "Дата источника": e.source_date or "без даты", "Срок подачи до": until, "Ссылка": e.url,
        "Тип источника": e.source_type, "Проверено через": VERIFIED_VIA.get(e.verified_via, e.verified_via),
    }


WIDTHS = {
    "№": 5, "Музей": 28, "Страна": 14, "Город": 16, "Что предлагают": 40, "Категория": 18, "Размеры": 18,
    "Кратко (рус.)": 60, "Статус предложения": 30, "Кто может получить": 40,
    "Доказательство: предлагают / отдают / продают": 70, "Ссылка на предложение": 40, "Дата источника": 12,
    "Доказательство: сняты / списаны / музей закрыт": 55, "Более ранние раунды (срок истёк)": 55,
    "Контактное лицо / отдел": 30, "Email": 30, "Телефон": 18, "Контактная страница": 40, "Следующий шаг": 50,
    "Проверено через": 26, "Что предлагали": 40, "Почему не в основном списке": 45, "Ссылка": 45,
    "Канал": 40, "Почему исключён": 70, "ID": 6, "Роль": 22, "Утверждение (из источника)": 70,
    "Текст источника": 50, "Срок подачи до": 16, "Тип источника": 14,
}


def _sheet(wb, title: str, fields: list[str], records: list[dict], fill: str) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(title)
    ws.append(fields)
    for r in records:
        ws.append([r.get(f, "") for f in fields])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    link_cols = {i for i, f in enumerate(fields) if f in ("Ссылка на предложение", "Ссылка", "Контактная страница")}
    for row in ws.iter_rows(min_row=2):
        for i, cell in enumerate(row):
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            v = str(cell.value or "")
            if i in link_cols and v.startswith("http") and "\n" not in v:
                cell.hyperlink = v
                cell.font = Font(color="0563C1", underline="single")
    for i, f in enumerate(fields, start=1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(f, 20)
    ws.freeze_panes = "C2" if "Музей" in fields[:2] else "B2"
    ws.auto_filter.ref = ws.dimensions


def write_verified_xlsx(files: list[MuseumFile], path: Path, today: dt.date | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    today = today or dt.date.today()
    rows, closed_rows, excluded, evidence_rows = verified_records(files, today)
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Отдают сейчас", VERIFIED_FIELDS, rows, "1E6B3A")
    _sheet(wb, "Срок истёк — спросить", CLOSED_FIELDS, closed_rows, "9C6500")
    _sheet(wb, "Исключено", EXCLUDED_FIELDS, excluded, "7F7F7F")
    _sheet(wb, "Доказательства", EVIDENCE_FIELDS, evidence_rows, "1F3A5F")

    ws = wb.create_sheet("Пояснения")
    notes = [
        f"Подтверждённые предложения музеев — {today.isoformat()}",
        "",
        "«Отдают сейчас»: только музеи, где источник прямо говорит, что предметы, которые музей больше "
        "не использует (сняты, заменены, списаны, музей закрыт), передаются, продаются или отдаются, "
        "и предложение не старше 2 лет, а срок объявления не истёк. Одна строка — один предмет/лот.",
        "«Срок истёк — спросить»: предлагали, но срок подачи истёк, данные старше 2 лет или сведения "
        "только из соцсетей. Имеет смысл спросить, что осталось.",
        "«Исключено»: каналы, которые продают/сдают в аренду новые выставки или доступны только "
        "организациям США — не подходят.",
        "«Доказательства»: каждое утверждение со ссылкой, датой и ролью.",
        "",
        "ВАЖНО: доказательства получены из сниппетов поисковой выдачи (страницы не открывались из среды "
        "исследования). Перед письмом откройте ссылку и убедитесь, что объявление на месте.",
        "Реконструкции и «экспонаты на складе, судьба неизвестна» сюда не входят — см. museum_finder.xlsx.",
    ]
    for line in notes:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 120
    for row in ws.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
