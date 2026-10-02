import csv

from mef.exporter import export_all
from mef.outreach import SUBJECT, build_draft
from mef.report import generate
from mef.rows import DATABASE_FIELDS


def test_outreach_draft_uses_facts_only(sample):
    d = build_draft(sample)
    assert d.subject == SUBJECT
    assert d.to == "jane@example-science.test"
    assert "Dear Jane Tester," in d.body
    assert "Forces Lab in 2024" in d.body
    assert "Bicycle generator" not in d.interest  # already transferred elsewhere


def test_export_has_required_fields(sample, tmp_path):
    paths = export_all([sample], out_dir=tmp_path, contacts_dir=tmp_path / "contacts")
    with paths["database"].open(encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    assert list(rows[0].keys()) == DATABASE_FIELDS
    assert rows[0]["Priority"] == "A"
    assert rows[0]["Old Exhibit"] == "Van de Graaff generator"
    assert paths["xlsx"].exists()


def test_report_sections(sample, tmp_path):
    run = generate([sample], label="test", out_root=tmp_path)
    text = (run / "REPORT.md").read_text(encoding="utf-8")
    for heading in ["1. EXHIBIT OPPORTUNITIES", "2. MUSEUM LEADS", "3. CONTACT LIST", "4. SOURCES",
                    "5. OUTREACH DRAFTS", "6. RESEARCH STATISTICS"]:
        assert heading in text
    assert (run / "outreach" / "example-science-centre.md").exists()
    assert "Foucault pendulum" in (run / "MUSEUM_CARDS.md").read_text(encoding="utf-8")


def test_candidates_workbook(sample, tmp_path):
    from openpyxl import load_workbook

    paths = export_all([sample], out_dir=tmp_path, contacts_dir=tmp_path / "contacts")
    wb = load_workbook(paths["candidates_xlsx"])
    assert wb.sheetnames == ["Кандидаты", "Сняты — спросить", "Каналы — запросить список", "Не рекомендовать",
                             "Исключённые каналы", "Правила"]
    ws = wb["Кандидаты"]
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    assert [r["Экспонат"] for r in rows] == ["Van de Graaff generator"]
    assert rows[0]["MUSEUM_RELEVANCE"] == "DIRECT_MATCH"
    assert rows[0]["Email"] == "jane@example-science.test"
    ask = [c.value for c in wb["Сняты — спросить"]["D"]][1:]
    assert ask == ["Foucault pendulum"]
