# Layout primitives for the Monad-style deck (python-pptx).
import copy, math, os, re, uuid
from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

# ---- palette (DESIGN.md, adapted) -------------------------------------------
PARCH, INK, GRAPH, SMOKE, ASH = 'F6F3F1', '242424', '4E4D4D', '797776', 'CECAC8'
BLUE, PERI = '2B59D1', 'CFDAF5'
CORAL, SKY, MINT, GOLD = 'FF9473', 'A0B5EB', 'A7FCCD', 'ECDA98'
PALETTE = {PARCH, INK, GRAPH, SMOKE, ASH, BLUE, PERI, CORAL, SKY, MINT, GOLD}

HEAD_FONT, BODY_FONT = 'Cambria', 'Consolas'
LANG = 'uz-Cyrl-UZ'

# ---- grid ---------------------------------------------------------------------
W, H, M = 13.333, 7.5, 0.6
CW = W - 2 * M
PILL_Y, PILL_H = 0.6, 0.3
TITLE_Y, TITLE_SIZE, TITLE_LEAD = 1.0, 36, 40
CT = 1.85                      # content top on every content slide
FB = H - M                     # footer baseline (bottom margin)
FN_W = CW - 1.3                # source/footnote width (page number sits right)
GAP, GAP_L = 0.3, 0.5

E = lambda v: Emu(int(round(v * 914400)))

# ---- text measurement (QA stand-in fonts: conservative Cambria, exact Consolas) --
_FONTS = {'head': os.path.expanduser('~/.fonts/QA-Cambria.ttf'), 'body': os.path.expanduser('~/.fonts/QA-Consolas.ttf')}
_fcache = {}
def tw(text, kind, size):
    f = _fcache.get(kind)
    if f is None:
        f = _fcache[kind] = ImageFont.truetype(_FONTS[kind], 400)
    return f.getlength(text) / 400 * size / 72.0

def wrap(text, kind, size, width):
    out = []
    for para in text.split('\n'):
        cur = ''
        for word in para.split(' '):
            t = (cur + ' ' + word) if cur else word
            if not cur or tw(t, kind, size) <= width:
                cur = t
            else:
                out.append(cur); cur = word
        out.append(cur)
    return out

def nlines(text, kind, size, width):
    return len(wrap(text, kind, size, width))

def lead_of(size, kind):
    if kind == 'head':
        return round(size * (1.0 if size >= 60 else 1.12), 1)
    return {9: 11.5, 10: 13, 11: 14, 12: 15.5, 14: 18, 16: 21}.get(size, round(size * 1.3, 1))

def text_h(text, kind, size, width, lead=None, after=0):
    lead = lead or lead_of(size, kind)
    paras = text.split('\n')
    n = sum(nlines(p, kind, size, width) for p in paras)
    return n * lead / 72.0 + after * (len(paras) - 1) / 72.0

WARN = []
def warn(msg):
    WARN.append(msg)

# ---- low-level XML helpers ------------------------------------------------------
def _strip_style(shape):
    st = shape._element.find(qn('p:style'))
    if st is not None:
        shape._element.remove(st)

def _no_effects(shape):
    spPr = shape._element.spPr
    if spPr.find(qn('a:effectLst')) is None:
        eff = etree.SubElement(spPr, qn('a:effectLst'))

def set_name(shape, name):
    if name:
        shape._element.xpath('./*[1]/p:cNvPr')[0].set('name', name)

def set_descr(shape, descr, decorative=False):
    cnv = shape._element.xpath('./*[1]/p:cNvPr')[0]
    cnv.set('descr', descr)
    if decorative:
        ext = etree.SubElement(etree.SubElement(cnv, qn('a:extLst')), qn('a:ext'))
        ext.set('uri', '{C183D7F6-B498-43B3-948B-1728B52AA6E4}')
        d = etree.SubElement(ext, '{http://schemas.microsoft.com/office/drawing/2017/decorative}decorative')
        d.set('val', '1')

# ---- runs / paragraphs --------------------------------------------------------
def _style_run(r, size, kind, color, spc=None):
    r.font.size = Pt(size)
    r.font.name = HEAD_FONT if kind == 'head' else BODY_FONT
    r.font.bold = False
    r.font.italic = False
    r.font.color.rgb = RGBColor.from_string(color)
    rPr = r._r.get_or_add_rPr()
    rPr.set('lang', LANG)
    if kind == 'head':
        rPr.set('spc', str(int(round(-0.02 * size * 100))))   # -2 % tracking
    elif spc is not None:
        rPr.set('spc', str(int(spc)))

def fill_tf(tf, paras, size=14, kind='body', color=INK, align='l', lead=None,
            after=0, upper=False, bullets=False):
    """paras: str (\\n = paragraph) or list of paragraphs; a paragraph is a str or a
    list of runs (text, {size,kind,color}) ."""
    if isinstance(paras, str):
        paras = paras.split('\n')
    al = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = al
        runs = [(p, {})] if isinstance(p, str) else p
        psize = max(o.get('size', size) for _, o in runs)
        pkind = runs[0][1].get('kind', kind)
        para.line_spacing = Pt(lead or lead_of(psize, pkind))
        if after and i < len(paras) - 1:
            para.space_after = Pt(after)
        if bullets:
            pPr = para._p.get_or_add_pPr()
            pPr.set('marL', str(int(0.2 * 914400))); pPr.set('indent', str(int(-0.2 * 914400)))
            bc = etree.SubElement(pPr, qn('a:buClr')); c = etree.SubElement(bc, qn('a:srgbClr')); c.set('val', SMOKE)
            bf = etree.SubElement(pPr, qn('a:buFont')); bf.set('typeface', BODY_FONT)
            bu = etree.SubElement(pPr, qn('a:buChar')); bu.set('char', '–')
        for t, o in runs:
            r = para.add_run()
            r.text = t.upper() if (o.get('upper', upper)) else t
            _style_run(r, o.get('size', size), o.get('kind', kind), o.get('color', color), o.get('spc'))

def tbox(sl, x, y, w, h, paras, size=14, kind='body', color=INK, align='l', anchor='t',
         lead=None, after=0, upper=False, name=None, check=True, bullets=False):
    shp = sl.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = shp.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    fill_tf(tf, paras, size, kind, color, align, lead, after, upper, bullets)
    set_name(shp, name)
    if check:
        plain = paras if isinstance(paras, str) else '\n'.join(
            p if isinstance(p, str) else ''.join(t for t, _ in p) for p in paras)
        if upper:
            plain = plain.upper()
        need = text_h(plain, kind, size, w - (0.2 if bullets else 0), lead, after)
        if need > h + 0.02:
            warn(f'overflow {name or plain[:40]!r}: need {need:.2f} > box {h:.2f}')
    return shp

# ---- shapes ----------------------------------------------------------------------
def rect(sl, x, y, w, h, fill=None, line=None, radius=None, lw=1.0, name=None):
    shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                              E(x), E(y), E(w), E(h))
    _strip_style(shp)
    if radius:
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    if fill:
        shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = RGBColor.from_string(line); shp.line.width = Pt(lw)
    else:
        shp.line.fill.background()
    _no_effects(shp)
    set_name(shp, name)
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return shp

def oval(sl, cx, cy, d, fill, line=None, name=None):
    shp = sl.shapes.add_shape(MSO_SHAPE.OVAL, E(cx - d / 2), E(cy - d / 2), E(d), E(d))
    _strip_style(shp)
    shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        shp.line.color.rgb = RGBColor.from_string(line); shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    _no_effects(shp)
    set_name(shp, name)
    return shp

def line(sl, x1, y1, x2, y2, color=ASH, lw=1.0, dash=False, name=None):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    _strip_style(c)
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(lw)
    if dash:
        ln = c.line._get_or_add_ln()
        pd = etree.SubElement(ln, qn('a:prstDash')); pd.set('val', 'dash')
    set_name(c, name)
    return c

def pill(sl, x, y, text, size=10, color=INK, border=ASH, fill=None, h=PILL_H, padx=0.16,
         w=None, align='c', name=None, upper=True):
    t = text.upper() if upper else text
    w = w or (tw(t, 'body', size) + 2 * padx)
    shp = rect(sl, x, y, w, h, fill=fill, line=border, radius=h / 2, name=name)
    tf = shp.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = E(padx) if align != 'c' else 0
    fill_tf(tf, [t], size=size, color=color, align=align, lead=lead_of(size, 'body'))
    return w

def picture(sl, path, x, y, w, h, radius=None, descr='', decorative=False, name=None, focus=(0.5, 0.5)):
    pic = sl.shapes.add_picture(path, E(x), E(y), E(w), E(h))
    iw, ih = Image.open(path).size
    fa, ia = w / h, iw / ih
    if ia > fa:      # crop sides
        keep = fa / ia; cut = 1 - keep
        pic.crop_left = cut * focus[0]; pic.crop_right = cut * (1 - focus[0])
    elif ia < fa:    # crop top/bottom
        keep = ia / fa; cut = 1 - keep
        pic.crop_top = cut * focus[1]; pic.crop_bottom = cut * (1 - focus[1])
    if radius:
        geom = pic._element.spPr.find(qn('a:prstGeom'))
        geom.set('prst', 'roundRect')
        av = geom.find(qn('a:avLst'))
        if av is None:
            av = etree.SubElement(geom, qn('a:avLst'))
        gd = etree.SubElement(av, qn('a:gd')); gd.set('name', 'adj')
        gd.set('fmla', f'val {int(radius / min(w, h) * 100000)}')
    set_descr(pic, descr, decorative)
    set_name(pic, name)
    return pic

def background(sl, color):
    sl.background.fill.solid()
    sl.background.fill.fore_color.rgb = RGBColor.from_string(color)

# ---- tables ------------------------------------------------------------------------
NOSTYLE = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'

def _ln(tag, color, w=12700):
    el = etree.Element(qn(tag)); el.set('w', str(w))
    if color:
        sf = etree.SubElement(el, qn('a:solidFill')); c = etree.SubElement(sf, qn('a:srgbClr')); c.set('val', color)
    else:
        etree.SubElement(el, qn('a:noFill'))
    return el

def table(sl, x, y, colw, rows, size=14, hdr=1, hdr_size=10, color=INK, hdr_color=GRAPH,
          rule=ASH, padx=0.08, pady=0.06, name=None, aligns=None, min_row=None):
    """rows: list of rows; a cell is str or dict(t, color, kind, size, fill, pattern, upper, align)."""
    nr, nc = len(rows), len(colw)
    norm = []
    heights = []
    for ri, row in enumerate(rows):
        nrow = []
        hmax = 0
        for ci, c in enumerate(row):
            d = {'t': c} if isinstance(c, str) else dict(c)
            ishdr = ri < hdr
            d.setdefault('size', hdr_size if ishdr else size)
            d.setdefault('kind', 'body')
            d.setdefault('color', hdr_color if ishdr else color)
            d.setdefault('upper', ishdr)
            d.setdefault('align', (aligns[ci] if aligns else 'l'))
            txt = d['t'].upper() if d['upper'] else d['t']
            hh = text_h(txt, d['kind'], d['size'], colw[ci] - 2 * padx) if txt else 0
            hmax = max(hmax, hh)
            nrow.append(d)
        rh = hmax + 2 * pady
        if min_row:
            rh = max(rh, min_row)
        heights.append(rh)
        norm.append(nrow)
    total = sum(heights)
    gf = sl.shapes.add_table(nr, nc, E(x), E(y), E(sum(colw)), E(total))
    set_name(gf, name)
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for a in ('firstRow', 'bandRow', 'firstCol', 'lastRow', 'lastCol', 'bandCol'):
        if a in tblPr.attrib:
            tblPr.set(a, '0')
    sid = tblPr.find(qn('a:tableStyleId'))
    if sid is None:
        sid = etree.SubElement(tblPr, qn('a:tableStyleId'))
    sid.text = NOSTYLE
    for ci, wv in enumerate(colw):
        tbl.columns[ci].width = E(wv)
    for ri, hv in enumerate(heights):
        tbl.rows[ri].height = E(hv)
    for ri, row in enumerate(norm):
        for ci, d in enumerate(row):
            cell = tbl.cell(ri, ci)
            tf = cell.text_frame
            tf.word_wrap = True
            fill_tf(tf, [d['t']] if d['t'] else [''], size=d['size'], kind=d['kind'], color=d['color'],
                    align=d['align'], upper=d['upper'])
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            for ch in list(tcPr):
                tcPr.remove(ch)
            for a, v in (('marL', padx), ('marR', padx), ('marT', pady), ('marB', pady)):
                tcPr.set(a, str(int(v * 914400)))
            tcPr.set('anchor', d.get('anchor', 't'))
            top_rule = rule if ri == 0 else None
            bottom_rule = rule
            tcPr.append(_ln('a:lnL', None)); tcPr.append(_ln('a:lnR', None))
            tcPr.append(_ln('a:lnT', d.get('top', top_rule if ri == 0 else None)))
            tcPr.append(_ln('a:lnB', d.get('bottom', bottom_rule)))
            if d.get('pattern'):
                pf = etree.SubElement(tcPr, qn('a:pattFill')); pf.set('prst', 'wdUpDiag')
                fg = etree.SubElement(etree.SubElement(pf, qn('a:fgClr')), qn('a:srgbClr')); fg.set('val', d['pattern'])
                bg = etree.SubElement(etree.SubElement(pf, qn('a:bgClr')), qn('a:srgbClr')); bg.set('val', d.get('fill', PARCH))
            elif d.get('fill'):
                sf = etree.SubElement(tcPr, qn('a:solidFill')); c = etree.SubElement(sf, qn('a:srgbClr')); c.set('val', d['fill'])
            else:
                etree.SubElement(tcPr, qn('a:noFill'))
    return total, heights
