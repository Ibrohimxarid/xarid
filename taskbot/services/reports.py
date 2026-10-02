"""Team dashboard, employee profile and Excel export."""

from __future__ import annotations

import io
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from taskbot.config import Settings
from taskbot.models import Task, TaskStatus, User
from taskbot.services import tasks as svc
from taskbot.services.dates import fmt_dt
from taskbot.services.kpi import EmployeeStats, Period, compute_stats, period_label, period_range_utc
from taskbot.services.scoring import fmt_num

STATUS_TITLES = {
    TaskStatus.PROPOSED: "Ожидает подтверждения",
    TaskStatus.ACTIVE: "В работе",
    TaskStatus.REWORK: "На доработке",
    TaskStatus.SUBMITTED: "На проверке",
    TaskStatus.DONE: "Оценена",
    TaskStatus.REJECTED: "Отклонена",
    TaskStatus.CANCELLED: "Отменена",
}


@dataclass
class TeamRow:
    employee: User
    stats: EmployeeStats
    tasks: list[Task] = field(default_factory=list)


@dataclass
class TeamReport:
    period: Period
    label: str
    rows: list[TeamRow]
    totals: EmployeeStats


async def team_report(session: AsyncSession, settings: Settings, period: Period, now: datetime) -> TeamReport:
    local_today = now.astimezone(settings.timezone).date()
    start, end = period_range_utc(period, local_today, settings.timezone)
    employees = await svc.active_employees(session)
    by_emp: dict[int, list[Task]] = defaultdict(list)
    for t in await svc.tasks_in_range(session, start, end):
        by_emp[t.employee_id].append(t)
    rows = [
        TeamRow(e, compute_stats(by_emp[e.id], now, settings.overdue_counts_as_zero), by_emp[e.id])
        for e in employees
    ]
    rows.sort(key=lambda r: (r.stats.coefficient is None, -(r.stats.coefficient or 0), r.employee.full_name))
    all_tasks = [t for r in rows for t in r.tasks]
    totals = compute_stats(all_tasks, now, settings.overdue_counts_as_zero)
    return TeamReport(period=period, label=period_label(period, local_today), rows=rows, totals=totals)


async def employee_stats(
    session: AsyncSession, settings: Settings, employee: User, period: Period, now: datetime
) -> EmployeeStats:
    local_today = now.astimezone(settings.timezone).date()
    start, end = period_range_utc(period, local_today, settings.timezone)
    tasks = await svc.tasks_in_range(session, start, end, employee.id)
    return compute_stats(tasks, now, settings.overdue_counts_as_zero)


def build_excel(report: TeamReport, settings: Settings) -> bytes:
    wb = Workbook()
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDEBF7")

    ws = wb.active
    ws.title = "Сводка"
    ws.append([f"Эффективность сотрудников — {report.label}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    headers = [
        "Сотрудник", "Коэффициент, %", "Задач", "Выполнено", "В работе", "На проверке",
        "Просрочено", "Выполнение в срок, %", "Сверх плана (>100%)", "Внесено самостоятельно",
    ]
    ws.append(headers)
    for row in report.rows:
        s = row.stats
        ws.append([
            row.employee.full_name,
            _round(s.coefficient),
            s.total, s.done, s.in_work, s.on_review, s.overdue,
            _round(s.on_time_pct), s.over_plan, s.self_added,
        ])
    _style_header(ws, 3, len(headers), header_font, header_fill)

    ws2 = wb.create_sheet("Задачи")
    headers2 = [
        "№", "Сотрудник", "Задача", "Ожидаемый результат", "План", "Срок", "Приоритет", "Вес, %",
        "Статус", "Сдано", "Фактический результат", "Факт", "Дополнительно", "Оценка AI, %",
        "Итоговая оценка, %", "Комментарий руководителя", "Просрочка",
    ]
    ws2.append(headers2)
    tz = settings.timezone
    for row in report.rows:
        for t in row.tasks:
            ws2.append([
                t.id, row.employee.full_name, t.title, t.expected_result,
                f"{fmt_num(t.plan_value)} {t.plan_unit or ''}".strip() if t.plan_value else "",
                fmt_dt(t.deadline, tz), t.priority, t.weight, STATUS_TITLES.get(t.status, t.status),
                fmt_dt(t.submitted_at, tz) if t.submitted_at else "", t.fact_text or "",
                t.fact_value if t.fact_value is not None else "", t.fact_extra or "",
                t.ai_score if t.ai_score is not None else "",
                t.final_score if t.final_score is not None else "",
                t.review_comment or "", "да" if t.missed_deadline else "",
            ])
    _style_header(ws2, 1, len(headers2), header_font, header_fill)
    for col, width in zip("ABCDEFGHIJKLMNOPQ", (6, 24, 40, 40, 14, 17, 10, 8, 16, 17, 40, 8, 30, 10, 10, 30, 10)):
        ws2.column_dimensions[col].width = width
    for r in ws2.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _round(v: float | None):
    return "" if v is None else round(v, 1)


def _style_header(ws, row: int, ncols: int, font: Font, fill: PatternFill) -> None:
    for i in range(1, ncols + 1):
        c = ws.cell(row=row, column=i)
        c.font = font
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
        if ws.title == "Сводка":
            ws.column_dimensions[get_column_letter(i)].width = 28 if i == 1 else 14
